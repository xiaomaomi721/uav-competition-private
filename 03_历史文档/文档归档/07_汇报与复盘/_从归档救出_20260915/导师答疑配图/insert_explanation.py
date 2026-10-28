#!/usr/bin/env python3
"""Insert the detailed '5597 vs 56000' explanation + data-flow diagram into the existing docx."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

BASE = Path(__file__).resolve().parent.parent
DOC = BASE / "无人机语义分割_实验方案与参数详解_导师答疑.docx"
FIG = BASE / "导师答疑配图" / "fig4_数据流讲解.png"

doc = Document(DOC)


def set_cn(run, name="宋体"):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def para(text="", bold=False, size=10.5, color=None, style=None, align=None):
    p = doc.add_paragraph()
    if style:
        p.style = doc.styles[style]
    r = p.add_run(text)
    r.font.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    set_cn(r)
    if align:
        p.alignment = align
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    r.font.size = Pt(10.5)
    set_cn(r)
    return p


def pic(path, width=6.3):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run()
    r.add_picture(str(path), width=Inches(width))
    return p


# 找到 1.3 标题段落作为插入锚点
anchor = None
for p in doc.paragraphs:
    if p.style.name.startswith("Heading") and "1.3" in p.text:
        anchor = p
        break
assert anchor is not None, "未找到 1.3 标题锚点"

# 构造要插入的段落（按顺序）
blocks = []

# 小标题
blocks.append(para("（划重点）56000 是「步数」，不是「图数」——一张图看懂", bold=True, size=12, color=RGBColor(0xD6, 0x27, 0x28)))

blocks.append(para("先说结论，避免继续混淆：我们手上并不是 5.6 万张图。官方给的全部训练图一共 6996 张，其中 5597 张拿来训练、1399 张拿来验证。56000 这个数字，是「训练步数」，跟「图数」是两个完全不同的东西。"))

blocks.append(para("为什么 6996 张要拆成 5597 + 1399 两份？", bold=True))
blocks.append(bullet("5597 张：用来「教」模型学东西（训练集）。"))
blocks.append(bullet("1399 张：不参与学习，专门当「考试卷」，检验模型学得好不好（验证集）。"))
blocks.append(para("就像考试不能拿讲过的原题来考自己——留一批模型没见过的图，才能测出它是不是真的学会了、能不能泛化到新图。"))

blocks.append(para("「反复喂 40 遍」是什么意思？", bold=True))
blocks.append(para("把 5597 张训练图从头到尾完整看一遍，叫 1 个 Epoch。因为显卡显存放不下全部图，一次只能喂一小批（一批 2 张，攒 2 批才算一次有效更新，等效一次 4 张），所以看完 5597 张需要约 1400 次参数更新，也就是 1 个 Epoch = 1400 步。反复看 40 遍 = 40 个 Epoch = 1400 × 40 = 56000 步。"))

blocks.append(para("所以 56000 的完整来历是：", bold=True))
blocks.append(para("5597 张训练图 ÷ 每步 4 张 ≈ 1400 步/Epoch，再 × 40 个 Epoch = 56000 步。每张训练图在整个训练过程中被反复看了约 40 次。"))

blocks.append(pic(FIG, 6.4))

blocks.append(para("一句话总结：6996 张图 → 5597 训练 + 1399 验证；5597 张看一遍 = 1 Epoch ≈ 1400 步；看 40 遍 = 56000 步。56000 是「步数」，不是「图数」。", bold=True, color=RGBColor(0x1F, 0x3B, 0x73)))

# 按顺序插入到锚点之前（addprevious 保持调用顺序）
for blk in blocks:
    anchor._p.addprevious(blk._p)

doc.save(DOC)
print("已插入讲解段落，段落总数：", len(doc.paragraphs))
